#!/usr/bin/env python3
"""S1636: Raimio Excel (VF/ZB lapai) -> s1636_x.php payload.
Naudojimas: python3 s1636_x_build.py tiekeju_prekiu_fiziniai_likuciai_S1632.xlsx
Stulpeliai lape: SKU | pavadinimas | eShoprent kiekis | MANO FIZINIS KIEKIS | savikaina | geriausia iki
"""
import sys,gzip,base64,re
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],data_only=True)
rows=[]
for lp in ('VF','ZB'):
    ws=wb[lp]
    for r in ws.iter_rows(min_row=2,values_only=True):
        if not r or not r[0]: continue
        sku=str(r[0]).strip(); fiz=r[3]; sav=r[4] if len(r)>4 else ''; gi=r[5] if len(r)>5 else ''
        if fiz in (None,''): continue
        gi=str(gi)[:10] if gi not in (None,'') else ''
        rows.append(f"{sku};{int(fiz)};{sav if sav not in (None,'') else ''};{gi};{lp}")
csv="\n".join(rows)
b64=base64.b64encode(gzip.compress(csv.encode())).decode()
t=open('s1636_x_sablonas.php').read()
assert t.count('__CSV__')==1
open('s1636_x.php','w').write(t.replace('__CSV__',b64))
print(f"eiluciu su fiziniu kiekiu: {len(rows)}; s1636_x.php paruostas")
